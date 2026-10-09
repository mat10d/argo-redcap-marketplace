#!/usr/bin/env python3
"""Field comments in a QA round — the three uses, pinned with known inputs and known outputs.

REDCap field comments are where RAs explain a blank, and sometimes where they typed the value
instead of entering it in the field. field_comments.py reads and judges them, for:

  1. Shorter worklists — a blank cell a comment already explains is painted blue ("confirm or
     ignore"), carries the comment as a cell note, and is not counted as a gap to fill. A returned
     blue cell the RA answered is read and checked like any other.
  2. Data stuck in comments — a comment that looks like the VALUE keeps the cell a gap and says
     "value may be in the comment — enter it in the field". The explanation lexicon wins.
  3. Comment contradicts value — a clear negation on a filled, non-"No", non-code value.

Without comments the builder behaves exactly as before. Every comment here is synthetic.
"""
from __future__ import annotations

import argparse
import csv
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
QA_SKILL = REPO / "plugins/argo-qa-specialist/skills/qa-worklists"
sys.path.insert(0, str(QA_SKILL))

try:
    import openpyxl
    import pandas  # noqa: F401
    import yaml  # noqa: F401
    DEPS = True
except ImportError:
    DEPS = False

import field_comments as fc  # noqa: E402
from qa_colours import AMBER_HEX, EXPLAINED_HEX, LEGACY_FLAG_HEXES, YELLOW_HEX  # noqa: E402

MDC4 = " | -666, Patient does not know | -777, Patient refused to answer | " \
       "-888, Missing in case notes | -999, Other missing"

DD = [
    {"field_name": "rid", "form_name": "f", "field_type": "text", "field_label": "Record"},
    {"field_name": "biopsy_result", "form_name": "f", "field_type": "radio",
     "field_label": "Biopsy result",
     "select_choices_or_calculations": "1, Benign | 2, Adenocarcinoma | "
                                       "3, Squamous cell carcinoma" + MDC4},
    {"field_name": "dx_date", "form_name": "f", "field_type": "text",
     "field_label": "Diagnosis date", "text_validation_type_or_show_slider_number": "date_dmy"},
    {"field_name": "weight", "form_name": "f", "field_type": "text", "field_label": "Weight",
     "text_validation_type_or_show_slider_number": "number", "text_validation_min": "20",
     "text_validation_max": "250"},
    {"field_name": "occupation", "form_name": "f", "field_type": "text",
     "field_label": "Occupation"},
    {"field_name": "surgery", "form_name": "f", "field_type": "yesno", "field_label": "Surgery"},
    {"field_name": "chemo", "form_name": "f", "field_type": "radio", "field_label": "Chemotherapy",
     "select_choices_or_calculations": "1, Yes | 0, No"},
    {"field_name": "followup", "form_name": "f", "field_type": "text", "field_label": "Follow-up",
     "branching_logic": "datediff([dx_date],'today','d')>30"},
]
FIELDS = ["biopsy_result", "dx_date", "weight", "occupation", "surgery", "chemo", "followup"]
COLS = ["rid", "redcap_data_access_group"] + FIELDS
DD_COLS = ["field_name", "form_name", "field_type", "field_label",
           "select_choices_or_calculations", "field_note",
           "text_validation_type_or_show_slider_number", "text_validation_min",
           "text_validation_max", "branching_logic", "matrix_group_name", "field_annotation"]
FILLED = {"biopsy_result": "1", "dx_date": "2024-01-01", "weight": "70",
          "occupation": "Teacher", "surgery": "0", "chemo": "0", "followup": "seen"}
# What differs from FILLED on each record: "" is a blank cell.
BEFORE = {
    "R1": {"biopsy_result": "", "surgery": "1"},
    "R2": {"biopsy_result": "", "weight": "-888"},
    "R3": {"biopsy_result": "", "followup": ""},
    "R4": {"dx_date": ""},
    "R5": {"dx_date": ""},
    "R6": {"weight": "", "occupation": ""},
    "R7": {"occupation": ""},
    "R8": {"surgery": "1", "occupation": ""},
    "R9": {"occupation": ""},
}
COMMENT_LOG = [  # Record, Field, User, Datetime, Comment — OAU 13.11.4's real headings
    ("R1", "biopsy_result", "adenocarcinoma"),                 # value: a choice label
    ("R1", "surgery", "No change, confirmed"),                 # NOT a disagreement (exception)
    ("R2", "biopsy_result", "Adenocarcinma"),                  # value: a choice label, one typo
    ("R2", "weight", "not weighed"),                           # on a code (-888): no disagreement
    ("R3", "biopsy_result", "Result not in chart"),            # explains the blank -> blue
    ("R3", "followup", "lost to follow up"),                   # amber stays amber, note added
    ("R4", "dx_date", "biopsy taken 12/03/2023"),              # value: a date
    ("R5", "dx_date", "Transferred on 12/03/2023"),            # lexicon wins over the date -> blue
    ("R6", "weight", "65 kg"),                                 # value: a number in range
    ("R6", "occupation", "Farmer"),                            # value: free text, no explanation
    ("R7", "occupation", "Unknown"),                           # explains the blank -> blue
    ("R8", "surgery", "No surgery done"),                      # disagrees: surgery = Yes
    ("R8", "chemo", "never had chemo"),                        # agrees: chemo = No
]

# The engineered with_MDC worklist: (record, column heading) -> fill.
B, Y, A = "blue", "yellow", "amber"
EXPECTED_FILLS = {
    ("R1", "Biopsy result"): Y, ("R2", "Biopsy result"): Y, ("R2", "Weight"): Y,
    ("R3", "Biopsy result"): B, ("R3", "Follow-up"): A,
    ("R4", "Diagnosis date"): Y, ("R5", "Diagnosis date"): B,
    ("R6", "Weight"): Y, ("R6", "Occupation"): Y, ("R7", "Occupation"): B,
    ("R8", "Occupation"): Y, ("R9", "Occupation"): Y,
}
VALUE_CELLS = {("R1", "biopsy_result"), ("R2", "biopsy_result"), ("R4", "dx_date"),
               ("R6", "weight"), ("R6", "occupation")}
DISAGREE_CELLS = {("R8", "surgery")}


def write_csv(path: Path, rows, cols):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


def rows_for(changes: dict) -> list:
    out = []
    for rid, diff in changes.items():
        r = dict(FILLED, rid=rid, redcap_data_access_group="site_alpha")
        r.update(diff)
        out.append(r)
    return out


def fill_kind(cell) -> str:
    rgb = str(cell.fill.fgColor.rgb or "").upper() if cell.fill and cell.fill.fgColor else ""
    for hexv, name in ((YELLOW_HEX, Y), (AMBER_HEX, A), (EXPLAINED_HEX, B)):
        if rgb.endswith(hexv):
            return name
    return ""


def build(tmp: Path, with_comments: bool) -> subprocess.CompletedProcess:
    cmd = [sys.executable, str(QA_SKILL / "build_worklists.py"),
           "--records-csv", str(tmp / "before.csv"), "--metadata-csv", str(tmp / "dd.csv"),
           "--fields", str(tmp / "qa_fields.yaml"), "--id-field", "rid", "--round=2026-10-01",
           "--out", str(tmp / ("with" if with_comments else "without"))]
    if with_comments:
        cmd += ["--comments", str(tmp / "comments.csv")]
    return subprocess.run(cmd, capture_output=True, text=True, timeout=300)


def sheet_cells(path: Path) -> dict:
    """{(record, heading): (fill, note text)} for every painted cell."""
    ws = openpyxl.load_workbook(path).active
    headers = [c.value for c in ws[1]]
    out = {}
    for r in range(3, ws.max_row + 1):
        for c in range(2, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            kind = fill_kind(cell)
            if kind:
                out[(ws.cell(row=r, column=1).value, headers[c - 1])] = (
                    kind, cell.comment.text if cell.comment else "")
    return out


@unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
class TestWorklistWithComments(unittest.TestCase):
    """Uses 1 and 2 in the builder, and the no-comments-unchanged rule."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = t = Path(tempfile.mkdtemp())
        write_csv(t / "dd.csv", DD, DD_COLS)
        write_csv(t / "before.csv", rows_for(BEFORE), COLS)
        (t / "qa_fields.yaml").write_text(
            "workbooks:\n  - name: clinical\n    title: Clinical\n    fields: ["
            + ", ".join(FIELDS) + "]\n")
        with open(t / "comments.csv", "w", newline="") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["Record", "Field", "User", "Datetime", "Comment"])
            for rec, field, text in COMMENT_LOG:
                w.writerow([rec, field, "ra_alpha", "2026-09-01 10:00:00", text])
        cls.with_ = build(t, True)
        cls.without = build(t, False)
        cls.book = t / "with/2026-10-01/with_MDC/clinical_site_alpha.xlsx"
        cls.plain = t / "without/2026-10-01/with_MDC/clinical_site_alpha.xlsx"

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_builds_complete(self):
        self.assertEqual(self.with_.returncode, 0, self.with_.stdout + self.with_.stderr)
        self.assertEqual(self.without.returncode, 0, self.without.stdout + self.without.stderr)

    def test_every_cell_gets_the_engineered_fill(self):
        got = {k: v[0] for k, v in sheet_cells(self.book).items()}
        self.assertEqual(got, EXPECTED_FILLS)

    def test_blue_cells_carry_the_comment_and_say_confirm_or_ignore(self):
        cells = sheet_cells(self.book)
        kind, note = cells[("R3", "Biopsy result")]
        self.assertIn("Result not in chart", note)
        self.assertIn("Confirm, or ignore", note)
        self.assertIn("Transferred on 12/03/2023", cells[("R5", "Diagnosis date")][1])

    def test_value_in_comment_stays_a_gap_and_says_so(self):
        cells = sheet_cells(self.book)
        for key in [("R1", "Biopsy result"), ("R2", "Biopsy result"), ("R4", "Diagnosis date"),
                    ("R6", "Weight"), ("R6", "Occupation")]:
            with self.subTest(key):
                kind, note = cells[key]
                self.assertEqual(kind, Y)
                self.assertIn("value may be in the comment", note)
                self.assertIn("enter it in the field", note)

    def test_the_cell_itself_stays_blank(self):
        """Blanks stay blank: a comment is never copied into a value."""
        ws = openpyxl.load_workbook(self.book).active
        headers = [c.value for c in ws[1]]
        col = headers.index("Occupation") + 1
        row = next(r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value == "R6")
        self.assertIn(ws.cell(row=row, column=col).value, (None, ""))

    def test_amber_is_not_turned_blue_by_a_comment(self):
        kind, note = sheet_cells(self.book)[("R3", "Follow-up")]
        self.assertEqual(kind, A)
        self.assertIn("lost to follow up", note)

    def test_run_summary_counts_the_explained_gaps(self):
        out = self.with_.stdout
        self.assertIn("Field comments: comments.csv (13 comment(s))", out)
        self.assertIn("3 gap(s) were already explained in REDCap", out)
        self.assertIn("9 gap(s) left to fill", out)
        self.assertIn("5 blank cell(s) where the value may be in the comment", out)
        self.assertIn("1 filled cell(s) where the comment and the value may disagree", out)
        self.assertIn("9 to fill, 3 already explained", out)

    def test_comment_checks_csv_lists_uses_2_and_3(self):
        with open(self.tmp / "with/2026-10-01/comment_checks.csv", newline="") as f:
            rows = list(csv.DictReader(f))
        got_value = {(r["Record"], r["Field"]) for r in rows
                     if r["Check"].startswith(fc.VALUE_IN_COMMENT)}
        got_disagree = {(r["Record"], r["Field"]) for r in rows if r["Check"] == fc.DISAGREES}
        self.assertEqual(got_value, VALUE_CELLS)
        self.assertEqual(got_disagree, DISAGREE_CELLS)
        r8 = next(r for r in rows if r["Record"] == "R8")
        self.assertEqual(r8["REDCap now"], "Yes")

    def test_without_comments_nothing_changes(self):
        cells = sheet_cells(self.plain)
        expected = {k: (Y if v == B else v) for k, v in EXPECTED_FILLS.items()}
        self.assertEqual({k: v[0] for k, v in cells.items()}, expected)
        self.assertFalse(any(note for _, note in cells.values()), "no notes without comments")
        self.assertNotIn("Field comments", self.without.stdout)
        self.assertNotIn("already explained", self.without.stdout)
        self.assertFalse((self.tmp / "without/2026-10-01/comment_checks.csv").exists())

    def test_blue_round_trips_through_a_returned_workbook(self):
        """The RA answers one blue cell and leaves the other two: one answer, nothing untouched."""
        import reconcile_return as rr
        import review_responses
        returned = self.tmp / "returned.xlsx"
        wb = openpyxl.load_workbook(self.book)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        row = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}
        ws.cell(row=row["R3"], column=headers.index("Biopsy result") + 1).value = "Benign"
        wb.save(returned)

        audit = review_responses.diff(str(self.book), str(returned))
        self.assertEqual(audit.by_record["R3"],
                         [review_responses.Answer("Biopsy result", "", "Benign", "explained")])
        blue = {k for k, (_, kind) in audit.flagged.items() if kind == "explained"}
        self.assertEqual(blue, {("R3", "Biopsy result"), ("R5", "Diagnosis date"),
                                ("R7", "Occupation")})

        after = dict(BEFORE, R3={"followup": ""})          # R3's biopsy result now entered: 1
        write_csv(self.tmp / "after.csv", rows_for(after), COLS)
        ns = argparse.Namespace(original=str(self.book), returned=str(returned), token_env=None,
                                url=None, records_csv=str(self.tmp / "after.csv"),
                                metadata_csv=str(self.tmp / "dd.csv"),
                                comments=str(self.tmp / "comments.csv"), comments_since="",
                                site=None, survey_forms="")
        rec = rr.gather(ns)[0]
        [check] = [c for c in rec.checks if c.record == "R3"]
        self.assertEqual((check.field, check.status, check.kind),
                         ("biopsy_result", "IN REDCAP", "explained"))
        self.assertEqual(rec.explained_left, 2)
        self.assertEqual(rec.untouched, 9, "the 9 to-fill gaps, not the blue ones")

        flags = {(f.record, f.field): f.reason for f in rec.comment_flags}
        self.assertEqual({k for k, v in flags.items() if v.startswith(fc.VALUE_IN_COMMENT)},
                         VALUE_CELLS)
        self.assertEqual({k for k, v in flags.items() if v == fc.DISAGREES}, DISAGREE_CELLS)

        report = rr.render(rec)
        self.assertIn("2 blue cell(s) — already explained by a field comment", report)
        self.assertIn("## Value may be in the comment — ask the RA to enter it in the field",
                      report)
        self.assertIn("## Comment and value may disagree — check", report)
        q = report.split("## For RA_questions.md", 1)[1]
        self.assertIn('### R6 — Occupation: the field is blank in REDCap, but its field comment '
                      'reads "Farmer".', q)
        self.assertNotIn("R8 — Surgery", q, "a disagreement is for the QA specialist to check")


class TestTheRules(unittest.TestCase):
    """The three judgements, one rule at a time."""

    META = {m["field_name"]: m for m in DD}

    def v(self, field, text):
        return fc.value_in_comment(text, self.META[field])

    def d(self, field, value, text):
        return fc.disagrees_with_value(text, value, self.META[field])

    def test_choice_labels_match_ignoring_case_spacing_and_small_typos(self):
        for text in ("Adenocarcinoma", "  adenocarcinoma. ", "Adenocarcinma",
                     "Squamous  cell carcinoma", "benign lesion"):
            with self.subTest(text):
                self.assertTrue(self.v("biopsy_result", text))
        for text in ("Biopsy sent to Lagos", "Ade", "-888"):
            with self.subTest(text):
                self.assertEqual(self.v("biopsy_result", text), "")

    def test_short_labels_need_an_exact_match(self):
        self.assertTrue(self.v("surgery", "yes"))
        self.assertEqual(self.v("surgery", "yea"), "")

    def test_dates(self):
        for text in ("2023-03-12", "seen 12.3.2023", "5 March 2023", "March 5, 2023"):
            with self.subTest(text):
                self.assertTrue(self.v("dx_date", text))
        for text in ("31/31/2023", "sometime in 2023", "08-08-8888"):
            with self.subTest(text):
                self.assertEqual(self.v("dx_date", text), "")

    def test_numbers_must_fit_the_field(self):
        self.assertTrue(self.v("weight", "65 kg"))
        self.assertEqual(self.v("weight", "300"), "", "above the field's maximum")
        self.assertEqual(self.v("weight", "-888"), "", "a missing-data code is not a value")
        self.assertEqual(self.v("weight", "weighed on 12/03/2023"), "", "a date is not a weight")

    def test_free_text_is_a_value_unless_it_explains(self):
        self.assertTrue(self.v("occupation", "Retired teacher"))
        self.assertEqual(self.v("occupation", "not stated in the folder"), "")

    def test_the_explanation_lexicon_wins(self):
        for field, text in (("biopsy_result", "Benign? result pending"),
                            ("dx_date", "Transferred on 12/03/2023"),
                            ("weight", "65 kg per referral letter, not available here"),
                            ("occupation", "Unknown"), ("biopsy_result", "patient died")):
            with self.subTest(text):
                self.assertTrue(fc.explains_blank(text))
                self.assertEqual(self.v(field, text), "")

    def test_a_bare_no_is_a_value_not_an_explanation(self):
        self.assertFalse(fc.explains_blank("No"))
        self.assertTrue(self.v("surgery", "No"))

    def test_disagreement_needs_a_clear_negation_at_the_start(self):
        for text in ("No surgery done", "Patient declined surgery", "did not have surgery",
                     "None", "never had an operation", "not done"):
            with self.subTest(text):
                self.assertEqual(self.d("surgery", "1", text), fc.DISAGREES)
        for text in ("Surgery done in 2023", "Operated, no complications",
                     "No change, confirmed", "no issues", "No surgery here but done at LUTH"):
            with self.subTest(text):
                self.assertEqual(self.d("surgery", "1", text), "")

    def test_no_disagreement_with_a_no_answer_or_a_code(self):
        self.assertEqual(self.d("surgery", "0", "No surgery done"), "")
        self.assertEqual(self.d("chemo", "0", "never had chemo"), "")
        self.assertEqual(self.d("weight", "-888", "not weighed"), "")
        self.assertEqual(self.d("weight", "0", "none"), "")
        self.assertEqual(self.d("dx_date", "8888-08-08", "no date"), "")
        self.assertEqual(self.d("surgery", "", "No surgery done"), "", "blank: rule 2's job")


class TestTheBlueColour(unittest.TestCase):
    def test_blue_is_its_own_colour(self):
        self.assertNotIn(EXPLAINED_HEX, (YELLOW_HEX, AMBER_HEX) + tuple(LEGACY_FLAG_HEXES))

    def test_the_reader_knows_it(self):
        import review_responses
        self.assertEqual(review_responses.FLAG_KINDS[EXPLAINED_HEX], "explained")


class TestCommentsThroughTheKey(unittest.TestCase):
    """The logging path, through a fake client — nothing can reach a real REDCap."""

    LOG = [
        {"timestamp": "2023-01-01 09:00", "username": "ra", "action": "Manage/Design ",
         "details": 'Add field comment (Record: R6, Field: occupation, Comment: "Farmer")'},
        {"timestamp": "2023-01-02 09:00", "username": "ra", "action": "Manage/Design ",
         "details": 'Add field comment (Record: R7, Field: occupation, Comment: "Unknown")'},
        {"timestamp": "2023-01-03 09:00", "username": "ra", "action": "Manage/Design ",
         "details": 'Delete field comment (Record: R7, Field: occupation, Comment: "Unknown")'},
    ]

    @unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
    def test_the_builder_reads_comments_from_logging_with_a_key(self):
        import build_worklists as bw      # puts the skill's vendored scripts/ on the path
        import argo_redcap_client
        seen, log = {}, self.LOG

        class FakeClient:
            def __init__(self, token, url=None, label=None):
                seen["url"] = url

            def _post(self, **params):
                seen.update(params)
                return log

        real = argo_redcap_client.RedcapClient
        argo_redcap_client.RedcapClient = FakeClient
        try:
            ns = argparse.Namespace(comments=None, token_env="SYN_TOKEN")
            comments, source = bw.load_comments(ns, "https://mock.argo.invalid/api/", "fake")
        finally:
            argo_redcap_client.RedcapClient = real
        self.assertEqual(seen["content"], "log")
        self.assertNotIn("beginTime", seen, "the whole history, not just this round")
        self.assertEqual([(c.record, c.field, c.text) for c in comments],
                         [("R6", "occupation", "Farmer")])
        self.assertIn("project logging", source)
        idx = fc.index_comments(comments, {m["field_name"]: m for m in DD})
        self.assertEqual(fc.value_in_comment(idx[("R6", "occupation")][0].text,
                                             DD[4]), fc.VALUE_IN_COMMENT)

    def test_no_logging_right_means_no_comments_and_no_error(self):
        class NoLog:
            def _post(self, **params):
                raise RuntimeError("403")
        self.assertEqual(fc.comments_through_key(NoLog()),
                         ([], "not available through the access key"))


class TestOneHomeForComments(unittest.TestCase):
    """Moved, not copied: the comment reader exists once, and reconcile_return re-exports it."""

    def test_defined_only_in_field_comments(self):
        for name in ("reconcile_return.py", "build_worklists.py", "upload_mdc.py"):
            text = (QA_SKILL / name).read_text()
            for defn in ("def read_comment_log", "def comments_from_logging",
                         "def comments_through_key", "def index_comments",
                         "FIELD_COMMENT_HEADERS = {"):
                with self.subTest(file=name, what=defn):
                    self.assertNotIn(defn, text)

    def test_reconcile_return_still_exposes_them(self):
        import reconcile_return as rr
        self.assertIs(rr.read_comment_log, fc.read_comment_log)
        self.assertIs(rr.Comment, fc.Comment)


if __name__ == "__main__":
    unittest.main(verbosity=2)

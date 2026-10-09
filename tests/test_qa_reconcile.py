#!/usr/bin/env python3
"""The QA check that REDCap really holds what the RA returned — and the one narrow write.

Decided 2026-10-09 (argo-core access-tiers.md): the QA specialist does not re-upload what the RAs
send back. They CHECK that REDCap now holds it (`reconcile_return.py`, read-only), and the only
thing they may upload is a missing-data code the RA returned for a cell that is still blank in
REDCap (`upload_mdc.py`). These tests pin both halves, and every refusal of the second.

Everything is synthetic and built in a temp folder: a tiny dictionary, a "before" export the
worklist is built from, an RA's returned workbook, an "after" export, a Field Comment Log, and a
file-backed mock REDCap (`.invalid` address — it cannot reach a real server).
"""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import site
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
QA_SKILL = REPO / "plugins/argo-qa-specialist/skills/qa-worklists"
CORE_REFS = REPO / "plugins/argo-core/skills/redcap-api/references"
sys.path.insert(0, str(QA_SKILL))

try:
    import openpyxl  # noqa: F401
    import pandas  # noqa: F401
    import yaml  # noqa: F401
    DEPS = True
except ImportError:
    DEPS = False

MDC4 = " | -666, Patient does not know | -777, Patient refused to answer | " \
       "-888, Missing in case notes | -999, Other missing"

DD = [
    {"field_name": "rid", "form_name": "f", "field_type": "text", "field_label": "Record"},
    {"field_name": "sex", "form_name": "f", "field_type": "radio", "field_label": "Sex",
     "select_choices_or_calculations": "1, Male | 2, Female" + MDC4},
    {"field_name": "dx_date", "form_name": "f", "field_type": "text",
     "field_label": "Diagnosis date", "text_validation_type_or_show_slider_number": "date_dmy",
     "field_note": "[06-06-6666, Patient does not know  08-08-8888, Missing in case notes  "
                   "09-09-9999, Other missing]"},
    {"field_name": "weight", "form_name": "f", "field_type": "text", "field_label": "Weight",
     "text_validation_type_or_show_slider_number": "number",
     "field_note": "[-666, Patient does not know  -777, Patient refused to answer  "
                   "-888, Missing in case notes  -999, Other missing]"},
    {"field_name": "symptoms", "form_name": "f", "field_type": "checkbox",
     "field_label": "Symptoms",
     "select_choices_or_calculations": "1, Pain | 2, Bleeding | -888, Missing in case notes"},
    {"field_name": "smoker", "form_name": "f", "field_type": "yesno", "field_label": "Smoker"},
    {"field_name": "occupation", "form_name": "f", "field_type": "text",
     "field_label": "Occupation"},
    {"field_name": "visit_dt", "form_name": "f", "field_type": "text", "field_label": "Visit time",
     "text_validation_type_or_show_slider_number": "datetime_dmy",
     "field_note": "[08-08-8888, Missing in case notes]"},
    {"field_name": "followup", "form_name": "f", "field_type": "text", "field_label": "Follow-up",
     "branching_logic": "datediff([dx_date],'today','d')>30"},
]
FIELDS = ["sex", "dx_date", "weight", "symptoms", "smoker", "occupation", "visit_dt", "followup"]
COLS = ["rid", "redcap_data_access_group", "sex", "dx_date", "weight", "symptoms___1",
        "symptoms___2", "symptoms____888", "smoker", "occupation", "visit_dt", "followup"]
IDS = ["R1", "R2", "R3", "R4", "R5", "R6"]


def before_rows():
    rows = []
    for rid in IDS:
        r = {c: "" for c in COLS}
        r.update(rid=rid, redcap_data_access_group="site_alpha",
                 symptoms___1="0", symptoms___2="0", symptoms____888="0")
        rows.append(r)
    rows[0]["weight"] = "60"          # R1's weight is not a gap: a comment on it is off-worklist
    return rows


def after_rows():
    """REDCap after the RA's round: some answers entered, some not, one entered differently."""
    rows = {r["rid"]: r for r in before_rows()}
    rows["R1"].update(sex="2", dx_date="2025-03-05", symptoms___1="1", symptoms___2="1",
                      occupation="Farmer")
    rows["R2"].update(weight="65", occupation="Trader")
    rows["R3"].update(occupation="Teacher")
    rows["R5"].update(sex="1")
    return list(rows.values())


# (record, column heading) -> what the RA typed into the returned workbook
RA_EDITS = {
    ("R1", "Sex"): "Female",                    # IN REDCAP (label vs code)
    ("R1", "Diagnosis date"): "05/03/2025",     # IN REDCAP (dd/mm vs YYYY-MM-DD)
    ("R1", "Symptoms"): "Pain, Bleeding",       # IN REDCAP (checkbox bits)
    ("R2", "Sex"): "Male",                      # NOT ENTERED
    ("R2", "Weight"): "70",                     # DIFFERS (REDCap 65)
    ("R2", "Occupation"): "filled",             # IN REDCAP ("filled" marker, REDCap has a value)
    ("R3", "Sex"): "-888",                      # NOT ENTERED, MDC -> uploadable
    ("R3", "Diagnosis date"): "08-08-8888",     # NOT ENTERED, date MDC -> uploadable
    ("R3", "Symptoms"): "Missing in case notes",  # NOT ENTERED, checkbox MDC -> uploadable
    ("R3", "Weight"): "-888",                   # NOT ENTERED, text MDC -> uploadable
    ("R3", "Smoker"): "-888",                   # UNCLEAR: yes/no has no -888
    ("R3", "Visit time"): "-888",               # NOT ENTERED, but datetime -> refused
    ("R3", "Occupation"): "-999",               # DIFFERS: REDCap holds "Teacher" -> refused
    ("R3", "Follow-up"): "-888",                # amber cell -> refused
    ("R4", "Sex"): "NO SURGERY",                # UNCLEAR
    ("R4", "Weight"): "-666",                   # NOT ENTERED, MDC, but the comment says died
}
NOTES = {"R5": "RESOLVED in REDCap", "R6": "patient died, chart not available"}

COMMENT_LOG = (
    "Record ID,Event,Field Name,Username,Comments,Timestamp\n"
    "R6,,dx_date,ra_alpha,Chart lost in the 2024 flood,2026-10-01 10:00\n"
    "R2,,sex,ra_alpha,Checked twice,2026-10-01 10:05\n"
    "R1,,weight,ra_alpha,Weighed on a broken scale,2026-10-01 10:10\n"
    "R4,,weight,ra_alpha,Patient died before the interview,2026-10-01 10:15\n"
)

EXPECTED = {
    ("R1", "sex"): "IN REDCAP", ("R1", "dx_date"): "IN REDCAP",
    ("R1", "symptoms"): "IN REDCAP",
    ("R2", "sex"): "NOT ENTERED", ("R2", "weight"): "DIFFERS", ("R2", "occupation"): "IN REDCAP",
    ("R3", "sex"): "NOT ENTERED", ("R3", "dx_date"): "NOT ENTERED",
    ("R3", "symptoms"): "NOT ENTERED", ("R3", "weight"): "NOT ENTERED",
    ("R3", "smoker"): "UNCLEAR", ("R3", "visit_dt"): "NOT ENTERED",
    ("R3", "occupation"): "DIFFERS", ("R3", "followup"): "NOT ENTERED",
    ("R4", "sex"): "UNCLEAR", ("R4", "weight"): "NOT ENTERED",
}


def write_csv(path: Path, rows, cols):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


DD_COLS = ["field_name", "form_name", "section_header", "field_type", "field_label",
           "select_choices_or_calculations", "field_note",
           "text_validation_type_or_show_slider_number", "text_validation_min",
           "text_validation_max", "identifier", "branching_logic", "required_field",
           "custom_alignment", "question_number", "matrix_group_name", "matrix_ranking",
           "field_annotation"]


class Round:
    """One synthetic round on disk: build -> return -> after-export -> comment log."""

    def __init__(self):
        self.tmp = Path(tempfile.mkdtemp())
        t = self.tmp
        write_csv(t / "dd.csv", DD, DD_COLS)
        write_csv(t / "before.csv", before_rows(), COLS)
        (t / "qa_fields.yaml").write_text(
            "workbooks:\n  - name: clinical\n    title: Clinical\n    fields: [" +
            ", ".join(FIELDS) + "]\n")
        proc = subprocess.run(
            [sys.executable, str(QA_SKILL / "build_worklists.py"), "--records-csv",
             str(t / "before.csv"), "--metadata-csv", str(t / "dd.csv"), "--fields",
             str(t / "qa_fields.yaml"), "--out", str(t / "worklists"), "--round=2026-10-01",
             "--id-field", "rid"], capture_output=True, text=True, timeout=300)
        assert proc.returncode == 0, proc.stdout + proc.stderr
        self.original = t / "worklists/2026-10-01/with_MDC/clinical_site_alpha.xlsx"
        self.returned = t / "returned.xlsx"
        wb = openpyxl.load_workbook(self.original)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        rows = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}
        for (rid, header), value in RA_EDITS.items():
            ws.cell(row=rows[rid], column=headers.index(header) + 1).value = value
        resp = headers.index("RESPONSE") + 1
        for rid, note in NOTES.items():
            ws.cell(row=rows[rid], column=resp).value = note
        wb.save(self.returned)
        write_csv(t / "after.csv", after_rows(), COLS)
        (t / "comments.csv").write_text(COMMENT_LOG)

    def cleanup(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


def args_for(r: Round, **kw):
    import argparse
    ns = argparse.Namespace(original=str(r.original), returned=str(r.returned), token_env=None,
                            url=None, records_csv=str(r.tmp / "after.csv"),
                            metadata_csv=str(r.tmp / "dd.csv"),
                            comments=str(r.tmp / "comments.csv"), comments_since="", site=None,
                            survey_forms="")
    for k, v in kw.items():
        setattr(ns, k, v)
    return ns


@unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
class TestReconcile(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import reconcile_return as rr
        cls.rr = rr
        cls.round = Round()
        cls.rec, _, cls.rows, cls.cols, cls.meta = rr.gather(args_for(cls.round))
        cls.report = rr.render(cls.rec)

    @classmethod
    def tearDownClass(cls):
        cls.round.cleanup()

    def status(self):
        return {(c.record, c.field): c.status for c in self.rec.checks if not c.from_note}

    def test_every_answered_cell_gets_the_engineered_status(self):
        self.assertEqual(self.status(), EXPECTED)

    def test_the_resolved_row_is_checked_cell_by_cell(self):
        noted = {(c.field): c.status for c in self.rec.checks if c.from_note and c.record == "R5"}
        self.assertEqual(noted["sex"], "IN REDCAP", "R5's sex was entered after the note")
        self.assertEqual(noted["dx_date"], "NOT ENTERED", "marked resolved, REDCap still blank")

    def test_a_note_that_explains_a_blank_is_not_a_not_entered(self):
        explained = {(c.record, c.field) for c in self.rec.still_blank_explained}
        self.assertIn(("R6", "dx_date"), explained)
        self.assertFalse(any(c.record == "R6" for c in self.rec.checks))

    def test_site_comes_from_the_records_dag(self):
        self.assertEqual(self.rec.site, "site_alpha")

    def test_report_leads_with_counts(self):
        head = self.report.split("## ", 1)[0]
        self.assertIn("| In REDCap | Not entered | Differs | Unclear |", head)
        n = {s: sum(1 for c in self.rec.checks if c.status == s) for s in self.rr.STATUSES}
        self.assertIn(f"| {n['IN REDCAP']} | {n['NOT ENTERED']} | {n['DIFFERS']} | "
                      f"{n['UNCLEAR']} |", head)
        self.assertIn("4 of the not-entered answers are missing-data codes", head)

    def test_report_lists_no_in_redcap_rows(self):
        """Counts first, then only what needs action — a matched cell is not an action."""
        self.assertNotIn("| R1 | Sex", self.report)

    def test_questions_block_is_ready_for_ra_questions_md(self):
        q = self.report.split("## For RA_questions.md", 1)[1]
        self.assertIn("\n## site_alpha\n", q)
        self.assertIn('### R2 — Sex: you wrote "Male" on the worklist, but REDCap is still blank.', q)
        self.assertIn("### R2 — Weight", q)
        self.assertIn("Which is right?", q)
        self.assertIn("### R4 — Sex", q)
        self.assertIn("### R5 — Diagnosis date: you marked this resolved", q)
        self.assertNotIn("### R3 — Sex", q, "an uploadable MDC is the QA's job, not an RA question")

    def test_the_questions_block_parses_as_one_site(self):
        import summarize_for_ra as s
        q = self.report.split("## For RA_questions.md", 1)[1].split("\n", 3)[3]
        path = self.round.tmp / "RA_questions.md"
        path.write_text(q)
        sections = s.parse_questions(str(path), warn=lambda *a: None)
        self.assertEqual(len(sections), 1)

    def test_field_comments_sit_beside_the_cells(self):
        self.assertIn("Checked twice", self.report)          # beside R2 sex (not entered)
        blank = self.report.split("Still blank, with a field comment", 1)[1].split("\n## ", 1)[0]
        self.assertIn("Chart lost in the 2024 flood", blank)
        self.assertIn("never converted", self.report)
        off = self.report.split("Field comments on cells not on the worklist", 1)[1]
        self.assertIn("Weighed on a broken scale", off.split("\n## ", 1)[0])

    def test_doubtful_mdcs_are_held_back_and_asked(self):
        held = {(c.record, c.field): why for c, why in self.rec.held_back}
        self.assertEqual(set(held), {("R4", "weight"), ("R3", "visit_dt")})
        self.assertIn("died", held[("R4", "weight")])
        self.assertIn("MDCs held back", self.report)
        q = self.report.split("## For RA_questions.md", 1)[1]
        self.assertIn("### R4 — Weight: you wrote -666 (Patient does not know). We held it back:", q)

    def test_a_high_mdc_share_is_flagged_once(self):
        self.assertEqual(self.report.count("High share of missing-data codes"), 1)

    def test_reconcile_never_writes(self):
        text = (QA_SKILL / "reconcile_return.py").read_text()
        for w in ("import_records", "overwriteBehavior", "import_metadata"):
            self.assertNotIn(w, text)


class TestMdcJudgement(unittest.TestCase):
    """An RA's missing-data code is checked, not rubber-stamped (2026-10-09)."""

    @classmethod
    def setUpClass(cls):
        import reconcile_return as rr
        cls.rr = rr
        cls.meta = {m["field_name"]: m for m in DD}

    def doubt(self, field, code, note="", comment="", surveys=()):
        c = self.rr.Check("R1", field, field, code, "blank", "NOT ENTERED", mdc=code)
        cs = [self.rr.Comment("R1", field, comment)] if comment else []
        return self.rr.mdc_doubt(c, self.meta[field], note, cs, surveys)

    def test_a_plain_code_with_a_fitting_note_passes(self):
        self.assertEqual(self.doubt("sex", "-888", note="not in the case notes"), "")
        self.assertEqual(self.doubt("weight", "-777", comment="patient refused"), "")
        self.assertEqual(self.doubt("sex", "-888"), "")

    def test_field_must_accept_the_code(self):
        self.assertIn("field note", self.doubt("occupation", "-888"))
        self.assertIn("not one of", self.doubt("symptoms", "-666"))
        self.assertIn("takes no", self.doubt("smoker", "-888"))
        self.assertIn("survey", self.doubt("sex", "-888", surveys=("f",)))
        m = dict(self.meta["sex"], matrix_group_name="phq9")
        c = self.rr.Check("R1", "sex", "sex", "-888", "blank", "NOT ENTERED", mdc="-888")
        self.assertIn("validated scale", self.rr.mdc_doubt(c, m, "", []))

    def test_999_needs_a_reason(self):
        self.assertIn("reason", self.doubt("sex", "-999"))
        self.assertEqual(self.doubt("sex", "-999", note="lab sample lost"), "")

    def test_code_must_fit_the_note(self):
        self.assertIn("-777", self.doubt("sex", "-888", note="the patient refused"))
        self.assertIn("died", self.doubt("weight", "-666", comment="Patient died in 2024"))

    def test_a_value_that_exists_is_a_question(self):
        self.assertIn("exists", self.doubt("sex", "-888", comment="it's in the paper chart"))
        self.assertEqual(self.doubt("sex", "-888", comment="not in the paper chart either"), "")

    def test_a_recode_over_an_existing_code_is_not_entered_and_never_uploadable(self):
        """with_MDC worklists flag cells already holding a code. REDCap still holding -999 when
        the RA wrote -888 means the RA didn't enter it — and the cell isn't blank to upload into."""
        c = self.rr.check_cell("R1", "Sex", "-888", "Other missing", {"sex": "-999"},
                               self.meta["sex"], "sex", ["sex"], "yellow")
        self.assertEqual((c.status, c.redcap_now, c.mdc), ("NOT ENTERED", "Other missing", "-888"))
        self.assertFalse(self.rr._mdc_candidate(c))

    def test_high_share_gets_one_line(self):
        C = lambda m: self.rr.Check("R", "f", "f", m or "x", "", "NOT ENTERED", mdc=m)  # noqa: E731
        self.assertIn("High share", self.rr.mdc_share_note([C("-888")] * 5 + [C("")] * 2))
        self.assertEqual(self.rr.mdc_share_note([C("-888")] * 4), "")
        self.assertEqual(self.rr.mdc_share_note([C("-888")] * 5 + [C("")] * 10), "")


class TestCommentLog(unittest.TestCase):
    def setUp(self):
        import reconcile_return as rr
        self.rr = rr
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_headings_are_matched_by_alias_not_position(self):
        p = self.tmp / "c.csv"
        p.write_text("Comment,Date/Time,User,Field,Record\nhello,2026-01-01,u,sex,R1\n")
        [c] = self.rr.read_comment_log(str(p))
        self.assertEqual((c.record, c.field, c.text, c.user), ("R1", "sex", "hello", "u"))

    def test_unknown_headings_fail_loudly_and_name_them(self):
        p = self.tmp / "c.csv"
        p.write_text("Patient,Question,Remark\nR1,sex,hello\n")
        with self.assertRaises(SystemExit) as cm:
            self.rr.read_comment_log(str(p))
        msg = str(cm.exception)
        for h in ("'Patient'", "'Question'", "'Remark'", "record, field, comment"):
            self.assertIn(h, msg)

    def test_logging_replay(self):
        E = lambda t, d, u="ra", a="Manage/Design ": {  # noqa: E731
            "timestamp": t, "username": u, "action": a, "details": d}
        entries = [
            E("2026-10-02 09:00", 'Edit field comment (Record: R1, Event: "Event 1", '
                                  'Field: "sex", Comment: "second, edited")'),
            E("2026-10-01 09:00", 'Add field comment (Record: R1, Event: "Event 1", '
                                  'Field: "sex", Comment: "first")'),
            E("2026-10-01 10:00", 'Add field comment (Record: R2, Field: "dx_date", '
                                  'Comment: "gone \\"lost\\"")'),
            E("2026-10-03 09:00", 'Delete field comment (Record: R2, Field: "dx_date", '
                                  'Comment: "gone \\"lost\\"")'),
            E("2026-10-01 11:00", 'Add field comment (Record: R3, Field: "x", Comment: "y")',
              a="Update record"),
        ]
        got = self.rr.comments_from_logging(entries)
        self.assertEqual([(c.record, c.field, c.text) for c in got],
                         [("R1", "sex", "second, edited")])

    def test_a_key_without_logging_rights_falls_back_quietly(self):
        class NoLog:
            def _post(self, **kw):
                raise RuntimeError("403")
        comments, why = self.rr.comments_through_key(NoLog(), "2026-10-01")
        self.assertEqual(comments, [])
        self.assertIn("not available", why)


@unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
class TestMdcUploadGate(unittest.TestCase):
    """What upload_mdc.py allows, and every refusal."""

    @classmethod
    def setUpClass(cls):
        import reconcile_return as rr
        import upload_mdc as um
        cls.um = um
        cls.round = Round()
        cls.rec, _, cls.rows, cls.cols, cls.meta = rr.gather(args_for(cls.round))
        cls.cells, cls.refused, cls.already = um.plan(cls.rec, cls.rows, cls.cols, cls.meta)

    @classmethod
    def tearDownClass(cls):
        cls.round.cleanup()

    def test_only_mdc_into_blank_flagged_cells(self):
        got = {(c.record, c.column, c.value) for c in self.cells}
        self.assertEqual(got, {("R3", "sex", "-888"), ("R3", "dx_date", "8888-08-08"),
                               ("R3", "symptoms____888", "1"), ("R3", "weight", "-888")})

    def refusal(self, rid, field):
        return next(why for c, why in self.refused if (c.record, c.field) == (rid, field))

    def test_a_held_back_code_is_refused(self):
        self.assertIn("held back", self.refusal("R4", "weight"))

    def test_a_real_value_is_refused(self):
        self.assertIn("not a missing-data code", self.refusal("R2", "sex"))
        self.assertIn("not a missing-data code", self.refusal("R4", "sex"))

    def test_a_non_blank_cell_is_refused(self):
        self.assertIn("already holds", self.refusal("R3", "occupation"))

    def test_an_amber_cell_is_refused(self):
        self.assertIn("amber", self.refusal("R3", "followup"))

    def test_a_datetime_field_is_refused(self):
        self.assertIn("date-and-time", self.refusal("R3", "visit_dt"))

    def test_yesno_cannot_take_a_code(self):
        self.assertIn("not a missing-data code", self.refusal("R3", "smoker"))
        self.assertIn("yesno", self.um.field_problem("smoker", {"field_type": "yesno"}, "rid"))

    def test_mdc_exempt_and_record_id_are_refused(self):
        self.assertIn("@MDC-EXEMPT", self.um.field_problem(
            "q1", {"field_type": "radio", "field_annotation": "@MDC-EXEMPT"}, "rid"))
        self.assertIn("record ID", self.um.field_problem("rid", {"field_type": "text"}, "rid"))

    def test_payload_is_one_row_per_record_blanks_elsewhere(self):
        text = self.um.payload_csv(self.cells, "rid")
        rows = list(csv.DictReader(io.StringIO(text)))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["rid"], "R3")
        self.um.check_payload(text, self.rec, self.rows, self.cols, self.meta)   # passes

    def assertRefused(self, text, needle):
        with self.assertRaises(self.um.MdcRefusal) as cm:
            self.um.check_payload(text, self.rec, self.rows, self.cols, self.meta)
        self.assertIn(needle, str(cm.exception))

    def test_final_gate_refuses_a_real_value(self):
        self.assertRefused("rid,sex\nR3,1\n", "not a missing-data code")

    def test_final_gate_refuses_a_non_blank_target(self):
        self.assertRefused("rid,occupation\nR3,-999\n", "already holds")

    def test_final_gate_refuses_a_cell_not_on_the_worklist(self):
        self.assertRefused("rid,weight\nR1,-888\n", "not on the worklist")

    def test_final_gate_refuses_a_new_record(self):
        self.assertRefused("rid,sex\nR99,-888\n", "isn't in REDCap")

    def test_final_gate_refuses_a_bad_checkbox_bit(self):
        self.assertRefused("rid,symptoms___1\nR3,1\n", "not a missing-data code")
        self.assertRefused("rid,symptoms____888\nR3,0\n", "not a missing-data code")

    def test_final_gate_refuses_a_display_form_date_and_datetime(self):
        self.assertRefused("rid,dx_date\nR3,08-08-8888\n", "not a missing-data code")
        self.assertRefused("rid,visit_dt\nR3,8888-08-08\n", "date-and-time")

    def test_final_gate_refuses_a_yesno(self):
        self.assertRefused("rid,smoker\nR3,-888\n", "can't hold")


MOCK_URL = "https://mock.argo.invalid/api/"


@unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
class TestMdcUploadThroughTheMock(unittest.TestCase):
    """The real command line, against the file-backed mock. Nothing can reach a real REDCap."""

    def setUp(self):
        sys.path.insert(0, str(REPO / "plugins/argo-core/skills/redcap-api/scripts"))
        import argo_redcap_mock as mock
        self.round = Round()
        t = self.round.tmp
        self.mock_dir = t / ".mock"
        p = self.mock_dir / "projects/study"
        p.mkdir(parents=True)
        tok = mock.fake_token("study")
        (self.mock_dir / "keys.json").write_text(json.dumps({tok: "study"}))
        (self.mock_dir / "config.json").write_text(json.dumps({"apply_writes": True}))
        (p / "project.json").write_text(json.dumps(
            {"project_id": "9077", "project_title": "SYN — Synthetic Cohort",
             "_stored_mode": "raw"}))
        (p / "metadata.json").write_text(json.dumps(DD))
        (p / "records.json").write_text(json.dumps(after_rows()))
        self.env_file = t / ".env"
        self.env_file.write_text(f"REDCAP_URL={MOCK_URL}\nARGO_REDCAP_MOCK=.mock\nCRC_TOKEN={tok}\n")
        self.home = t / "home"
        self.home.mkdir()

    def tearDown(self):
        self.round.cleanup()

    def run_cli(self, *extra):
        env = {k: v for k, v in os.environ.items()
               if k not in ("REDCAP_URL", "ARGO_REDCAP_MOCK", "CRC_TOKEN")
               and not k.endswith("_TOKEN") and not k.endswith("_REQUEST")}
        env.update(ARGO_ENV_FILE=str(self.env_file), HOME=str(self.home),
                   PYTHONUSERBASE=site.getuserbase())
        return subprocess.run(
            [sys.executable, str(QA_SKILL / "upload_mdc.py"), str(self.round.original),
             str(self.round.returned), "--token-env", "CRC_TOKEN",
             "--comments", str(self.round.tmp / "comments.csv"), *extra],
            capture_output=True, text=True, timeout=300, env=env, cwd=str(self.round.tmp))

    def writes(self):
        f = self.mock_dir / "WRITES.jsonl"
        return [json.loads(l) for l in f.read_text().splitlines()] if f.exists() else []

    def test_upload_without_a_preview_is_refused(self):
        p = self.run_cli("--upload", "--expect-project", "9077", "--snapshot-dir", "snaps")
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("REFUSING", p.stdout + p.stderr)
        self.assertEqual(self.writes(), [])

    def test_dry_run_then_upload_sends_only_the_codes_and_reads_them_back(self):
        p = self.run_cli("--dry-run")
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        self.assertIn("Nothing has been changed", p.stdout)
        self.assertEqual(self.writes(), [])
        p = self.run_cli("--upload", "--expect-project", "9077", "--snapshot-dir", "snaps")
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        [w] = self.writes()
        self.assertEqual(w["overwriteBehavior"], "normal")
        sent = {(r["rid"], k, v) for r in w["data"] for k, v in r.items() if k != "rid" and v}
        self.assertEqual(sent, {("R3", "sex", "-888"), ("R3", "dx_date", "8888-08-08"),
                                ("R3", "symptoms____888", "1"), ("R3", "weight", "-888")})
        self.assertIn("All 4 code(s) are in REDCap", p.stdout)
        self.assertTrue(list((self.round.tmp / "snaps").glob("snapshot_*_pre-mdc.csv")))

    def test_the_wrong_project_is_refused(self):
        self.run_cli("--dry-run")
        p = self.run_cli("--upload", "--expect-project", "1234", "--snapshot-dir", "snaps")
        self.assertNotEqual(p.returncode, 0)
        self.assertEqual(self.writes(), [])

    def test_no_project_named_is_refused(self):
        self.run_cli("--dry-run")
        p = self.run_cli("--upload", "--snapshot-dir", "snaps")
        self.assertNotEqual(p.returncode, 0)
        self.assertIn("name the project", p.stdout + p.stderr)
        self.assertEqual(self.writes(), [])

    def test_reconcile_through_the_key_is_read_only(self):
        env = {k: v for k, v in os.environ.items()
               if k not in ("REDCAP_URL", "ARGO_REDCAP_MOCK", "CRC_TOKEN")}
        env.update(ARGO_ENV_FILE=str(self.env_file), HOME=str(self.home),
                   PYTHONUSERBASE=site.getuserbase())
        p = subprocess.run(
            [sys.executable, str(QA_SKILL / "reconcile_return.py"), str(self.round.original),
             str(self.round.returned), "--token-env", "CRC_TOKEN"],
            capture_output=True, text=True, timeout=300, env=env, cwd=str(self.round.tmp))
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        self.assertIn("# REDCap check — site_alpha", p.stdout)
        self.assertIn("Field comments: not included", p.stdout,
                      "the mock has no logging, so the key path falls back quietly")
        self.assertEqual(self.writes(), [])


@unittest.skipIf(not DEPS, "pandas/openpyxl/yaml not installed")
class TestMdcImportFile(unittest.TestCase):
    def test_no_key_writes_a_file_without_checkbox_columns(self):
        r = Round()
        try:
            out = r.tmp / "mdc_import.csv"
            p = subprocess.run(
                [sys.executable, str(QA_SKILL / "upload_mdc.py"), str(r.original), str(r.returned),
                 "--records-csv", str(r.tmp / "after.csv"), "--metadata-csv", str(r.tmp / "dd.csv"),
                 "--comments", str(r.tmp / "comments.csv"), "--import-file", str(out)], capture_output=True, text=True, timeout=300,
                cwd=str(r.tmp))
            self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
            rows = list(csv.DictReader(out.open()))
            self.assertEqual(rows, [{"rid": "R3", "sex": "-888", "dx_date": "8888-08-08",
                                     "weight": "-888"}])
            self.assertIn("the RA ticks it in REDCap", p.stdout)
        finally:
            r.cleanup()


class TestDoctrineIsConsistent(unittest.TestCase):
    """2026-10-09: no re-upload by QA, except missing-data codes into blank cells. Every place
    that states the QA write rule must state THIS rule — an older 'never writes' left in one file
    is how a session follows the wrong one."""

    SKILL = (QA_SKILL / "SKILL.md").read_text()
    TIERS = (CORE_REFS / "access-tiers.md").read_text()
    GOTCHAS = (CORE_REFS / "redcap-api-gotchas.md").read_text()

    def test_access_tiers_records_the_decision_dated(self):
        self.assertIn("2026-10-09", self.TIERS)
        self.assertIn("upload_mdc.py", self.TIERS)

    def test_gotchas_policy_names_the_exception(self):
        sec0 = self.GOTCHAS.split("## 0.", 1)[1].split("\n## 1.", 1)[0]
        self.assertIn("upload_mdc.py", sec0)
        self.assertIn("2026-10-09", sec0)

    def test_skill_task_2_is_built_on_the_direct_check(self):
        task2 = self.SKILL.split("## Task 2", 1)[1]
        self.assertIn("reconcile_return.py", task2)
        self.assertIn("upload_mdc.py", task2)
        self.assertIn("--comments", task2)
        for status in ("IN REDCAP", "NOT ENTERED", "DIFFERS"):
            self.assertIn(status, task2)

    def test_no_file_still_says_the_qa_round_never_writes(self):
        stale = ("a QA round never writes back", "The QA loop is read-only",
                 "the QA/analysis loop is read-only", "No write-back (by design")
        for path in [QA_SKILL / "SKILL.md", QA_SKILL / "references/migration-push.md",
                     QA_SKILL / "push_updates.py", QA_SKILL / "verify_push.py",
                     CORE_REFS / "redcap-api-gotchas.md", CORE_REFS / "access-tiers.md"]:
            text = path.read_text()
            for s in stale:
                with self.subTest(path=path.name, phrase=s):
                    self.assertNotIn(s, text)


if __name__ == "__main__":
    unittest.main()
